"""Tailor a resume from an approved .docx template, then export docx + PDF.

Design rule: the .docx files in `resumes/` are the ONLY source of layout. We
never rebuild a document — we open a copy and rewrite text in place, so
margins / fonts / tab stops / section rules stay exactly as tuned in Word.

Three ways to change a resume, matching how the work actually shows up:

  3. Job fits an existing template  (the common case — most general hires)
        --worksheet  w.json                 dump every bullet as editable JSON
        (edit w.json: reorder / rewrite / add / drop bullets, skills, tagline)
        --apply w.json --out X.docx --pdf   write the edited copy

  2. Job is domain-specific
        same worksheet round-trip, just a bigger edit.

  1. A role type keeps recurring
        --apply w.json --out X.docx --pdf --save-template mle_speech
        promotes the result into resumes/ and registers it in
        resumes/variants.json so it becomes a baseline like sde_ng.

Other commands:
    --show                    print structure (entry keys + bullets)
    --list                    list registered variants
    --jd FILE                 with --worksheet: annotate JD keywords the
                              resume is missing, to guide the edit

Usage:
    python3 scripts/resume_build.py --variant sde_mle_ng --worksheet /tmp/w.json
    python3 scripts/resume_build.py --apply /tmp/w.json \
        --out application_records/<date>_<company>_<role>/resume.docx --pdf
"""

from __future__ import annotations

import argparse
import copy
import json
import re
import subprocess
import sys
from pathlib import Path

import docx

_ROOT = Path(__file__).resolve().parents[1]
_RESUMES = _ROOT / "resumes"
_REGISTRY = _RESUMES / "variants.json"
_CLAIMABLE = _RESUMES / "claimable_terms.json"

# Resume filenames come from config/profile.yaml so this file carries no
# assumption about whose resume it is.
import yaml as _yaml  # noqa: E402
with (_ROOT / "config" / "profile.yaml").open(encoding="utf-8") as _pf:
    _RESUME_CFG = _yaml.safe_load(_pf).get("resumes", {})
_STEM = _RESUME_CFG.get("stem", "Resume")
BUILTIN_VARIANTS = {k: f"{_STEM}_{v}.docx"
                    for k, v in _RESUME_CFG.get("variants", {}).items()}

SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"


def variants() -> dict:
    v = dict(BUILTIN_VARIANTS)
    if _REGISTRY.exists():
        v.update(json.loads(_REGISTRY.read_text()))
    return v


# ---------- structure ----------

def _is_heading(p) -> bool:
    t = p.text.strip()
    return bool(t) and t.isupper() and len(p.runs) == 1 and p.runs[0].bold


def _is_bullet(p) -> bool:
    return p.style is not None and p.style.name == "List Paragraph"


def _bold_prefix(p) -> str:
    out = ""
    for r in p.runs:
        if r.bold:
            out += r.text
        else:
            break
    return out.strip()


def _split_date(text: str) -> tuple[str, str]:
    parts = re.split(r"\t+|\s{6,}", text)
    if len(parts) >= 2 and re.search(r"(19|20)\d\d", parts[-1]):
        return " ".join(x.strip() for x in parts[:-1]).strip(), parts[-1].strip()
    return text.strip(), ""


def index(doc) -> dict:
    """entry-key -> {head, bullets, section}; skills label -> paragraph."""
    entries, skills, order = {}, {}, []
    section, cur = None, None
    for p in doc.paragraphs[3:]:
        if not p.text.strip():
            continue
        if _is_heading(p):
            section, cur = p.text.strip(), None
            continue
        if _is_bullet(p):
            if cur:
                entries[cur]["bullets"].append(p)
            continue
        if section and section.startswith("TECHNICAL") and ":" in p.text:
            skills[p.text.split(":", 1)[0].strip()] = p
            continue
        key = _bold_prefix(p)
        if key:
            entries[key] = {"head": p, "bullets": [], "section": section}
            order.append(key)
            cur = key
    return {"entries": entries, "skills": skills, "order": order}


# ---------- in-place text rewriting ----------

def set_text(p, text: str) -> None:
    """Rewrite a paragraph keeping the first run's character formatting."""
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r._element.getparent().remove(r._element)


def set_bullets(entry: dict, new: list[str]) -> None:
    olds = entry["bullets"]
    if not olds:
        raise ValueError("entry has no bullet to clone formatting from")
    n, m = len(new), len(olds)
    for i in range(min(n, m)):
        set_text(olds[i], new[i])
    if n > m:                                    # clone last bullet's XML
        anchor = olds[m - 1]._p
        for txt in new[m:]:
            el = copy.deepcopy(olds[m - 1]._p)
            anchor.addnext(el)
            anchor = el
            para = docx.text.paragraph.Paragraph(el, olds[m - 1]._parent)
            set_text(para, txt)
            entry["bullets"].append(para)
    elif n < m:
        for p in olds[n:]:
            p._p.getparent().remove(p._p)
        entry["bullets"] = olds[:n]


def set_skill(p, value: str, label: str | None = None) -> None:
    runs = p.runs
    if label is not None and runs:
        runs[0].text = f"{label}: "
    if len(runs) >= 2:
        runs[1].text = value
        for r in runs[2:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(value).bold = False


def set_heading(p, left: str) -> None:
    """Heading line is `<bold>Role</bold><plain> — Company</plain><tab><date>`.
    Only the pre-tab part is rewritten; the tab + date runs are left alone."""
    runs = p.runs
    if not runs:
        return
    bold, plain = (left.split(" — ", 1) + [""])[:2]
    tab_i = next((i for i, r in enumerate(runs) if "\t" in r.text), None)
    body = runs[:tab_i] if tab_i is not None else runs
    if not body:
        return
    body[0].text = bold
    if len(body) > 1:
        body[1].text = (" — " + plain) if plain else ""
        for r in body[2:]:
            r.text = ""
    elif plain:
        body[0].text = left


# ---------- worksheet ----------

def to_worksheet(doc, variant: str) -> dict:
    idx = index(doc)
    sections, seen = [], {}
    for key in idx["order"]:
        e = idx["entries"][key]
        sec = seen.get(e["section"])
        if sec is None:
            sec = {"title": e["section"], "entries": []}
            seen[e["section"]] = sec
            sections.append(sec)
        left, date = _split_date(e["head"].text)
        sec["entries"].append({
            "key": key, "heading": left, "date_readonly": date,
            "bullets": [b.text for b in e["bullets"]],
        })
    return {
        "variant": variant,
        "tagline": doc.paragraphs[2].text.strip(),
        "sections": sections,
        "skills": {k: p.text.split(":", 1)[1].strip() for k, p in idx["skills"].items()},
    }


def apply_worksheet(doc, ws: dict) -> list[str]:
    idx, log = index(doc), []

    if ws.get("tagline") and ws["tagline"] != doc.paragraphs[2].text.strip():
        set_text(doc.paragraphs[2], ws["tagline"])
        log.append("tagline rewritten")

    for sec in ws.get("sections", []):
        for ent in sec["entries"]:
            key = ent["key"]
            if key not in idx["entries"]:
                raise KeyError(f"entry {key!r} not in template "
                               f"(have: {list(idx['entries'])})")
            e = idx["entries"][key]
            old_head, _ = _split_date(e["head"].text)
            if ent.get("heading") and ent["heading"] != old_head:
                set_heading(e["head"], ent["heading"])
                log.append(f"{key}: heading -> {ent['heading']}")
            old = [b.text for b in e["bullets"]]
            new = ent.get("bullets", old)
            if new != old:
                set_bullets(e, new)
                changed = sum(1 for a, b in zip(old, new) if a != b)
                log.append(f"{key}: {len(old)} -> {len(new)} bullets "
                           f"({changed} reworded/reordered in place)")

    for label, spec in (ws.get("skills") or {}).items():
        if label not in idx["skills"]:
            raise KeyError(f"skills label {label!r} not in template "
                           f"(have: {list(idx['skills'])})")
        p = idx["skills"][label]
        # a plain string rewrites the value; {"label","value"} also renames it
        new_label, value = (None, spec) if isinstance(spec, str) else \
            (spec.get("label"), spec["value"])
        cur = p.text.split(":", 1)[1].strip()
        if cur != value or (new_label and new_label != label):
            set_skill(p, value, new_label)
            log.append(f"skills[{label}] rewritten"
                       + (f" -> relabelled {new_label!r}" if new_label else ""))

    return log


# ---------- JD keyword diff ----------

# A curated ATS vocabulary. We only report terms from this list, because raw
# frequency ranking surfaces noise ("google", "san", "take") that no resume
# should chase. Multi-word terms are matched as phrases.
ATS_TERMS = [
    # languages / core CS
    "python", "c++", "java", "javascript", "typescript", "go", "golang", "rust",
    "scala", "kotlin", "swift", "sql", "bash", "shell scripting", "matlab", "r",
    "data structures", "algorithms", "object-oriented", "design patterns",
    "complexity analysis", "concurrency", "multithreading", "parallel systems",
    "compilers", "operating systems", "computer architecture",
    # systems / infra
    "distributed systems", "distributed computing", "large-scale", "scalability",
    "high availability", "fault tolerance", "microservices", "system design",
    "large software systems", "backend", "frontend", "full-stack", "web development",
    "mobile development", "android", "ios", "api", "rest", "grpc", "graphql",
    "unix", "linux", "docker", "kubernetes", "terraform", "aws", "gcp", "azure",
    "ci/cd", "devops", "sre", "site reliability", "observability", "monitoring",
    "prometheus", "load testing", "performance optimization", "profiling",
    "caching", "networking", "data storage", "databases", "postgresql", "mysql",
    "nosql", "spark", "hadoop", "kafka", "airflow", "etl", "data pipelines",
    "message queue", "security", "cryptography", "authentication",
    # ml / ai
    "machine learning", "deep learning", "neural networks", "pytorch",
    "tensorflow", "hugging face", "scikit-learn", "nlp",
    "natural language processing", "information retrieval", "search", "ranking",
    "recommendation", "computer vision", "speech recognition", "asr",
    "reinforcement learning", "llm", "large language model", "generative ai",
    "genai", "rag", "retrieval-augmented", "embeddings", "vector database",
    "fine-tuning", "lora", "quantization", "inference", "model serving",
    "distributed training", "ddp", "fsdp", "cuda", "gpu", "onnx", "vllm",
    "mlops", "feature engineering", "a/b testing", "experimentation",
    "ai productivity tools", "ai tools", "agents", "agentic", "tool calling",
    "prompt engineering", "evaluation", "benchmarking",
    # practice / soft-technical
    "code review", "unit testing", "integration testing", "test-driven",
    "debugging", "documentation", "version control", "git", "agile",
    "cross-functional", "technical leadership", "mentoring", "open-source",
    "accessible technologies", "accessibility", "ui design", "user experience",
    "deployment", "rollback", "incident response", "on-call",
]


def _present(term: str, text: str) -> bool:
    """Word-boundary-ish containment that tolerates +, /, - in terms and an
    optional trailing 's' on the last word (JD says "tools", resume says "tool")."""
    t = re.escape(term).replace(r"\ ", r"[\s-]+")
    t = re.sub(r"s$", "", t) + "s?"
    return re.search(rf"(?<![a-z0-9]){t}(?![a-z0-9])", text) is not None


def jd_keywords(jd: str, top: int = 0) -> list[str]:
    """ATS terms the JD actually uses, ordered by how often it repeats them."""
    low = jd.lower()
    hits = [(t, len(re.findall(re.escape(t).replace(r"\ ", r"[\s-]+"), low)))
            for t in ATS_TERMS]
    hits = [(t, n) for t, n in hits if n and _present(t, low)]
    hits.sort(key=lambda x: (-x[1], ATS_TERMS.index(x[0])))
    return [t for t, _ in hits]


def resume_text(ws: dict) -> str:
    parts = [ws.get("tagline", "")]
    for s in ws.get("sections", []):
        for e in s["entries"]:
            parts.append(e.get("heading", ""))
            parts.extend(e.get("bullets", []))
    for spec in (ws.get("skills") or {}).values():
        if isinstance(spec, str):
            parts.append(spec)
        else:                                  # {"label": ..., "value": ...}
            parts.append(spec.get("label", ""))
            parts.append(spec.get("value", ""))
    return " ".join(parts).lower()


# ---------- pdf ----------

def to_pdf(docx_path: Path) -> Path:
    if not Path(SOFFICE).exists():
        raise RuntimeError(f"LibreOffice not found at {SOFFICE}")
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf",
                    "--outdir", str(docx_path.parent), str(docx_path)],
                   check=True, capture_output=True, timeout=180)
    pdf = docx_path.parent / (docx_path.stem + ".pdf")
    if not pdf.exists():
        raise RuntimeError("conversion produced no pdf")
    return pdf


def page_count(pdf: Path) -> int:
    return len(re.findall(rb"/Type\s*/Page[^s]", pdf.read_bytes()))


# ---------- cli ----------

def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--variant", help="template name (see --list)")
    ap.add_argument("--list", action="store_true")
    ap.add_argument("--show", action="store_true")
    ap.add_argument("--worksheet", help="write editable worksheet JSON here")
    ap.add_argument("--jd", help="JD text file; with --worksheet, annotate missing keywords")
    ap.add_argument("--apply", help="worksheet JSON to apply")
    ap.add_argument("--out", help="output .docx path")
    ap.add_argument("--pdf", action="store_true")
    ap.add_argument("--save-template", metavar="NAME",
                    help="promote the result into resumes/ as a new variant. "
                         "Per rulebook 3.1a this needs cross-JD evidence and the "
                         "user's OK; --reason is mandatory and gets logged.")
    ap.add_argument("--reason", help="why this template change is justified "
                                     "(required with --save-template)")
    args = ap.parse_args()

    V = variants()

    if args.list:
        for k, f in V.items():
            mark = "" if (_RESUMES / f).exists() else "   [MISSING FILE]"
            print(f"  {k:16s} {f}{mark}")
        return 0

    if args.apply:
        ws = json.loads(Path(args.apply).read_text())
        variant = args.variant or ws.get("variant")
    else:
        variant = args.variant
    if not variant:
        print("--variant required (or a worksheet carrying one)"); return 2
    if variant not in V:
        print(f"unknown variant {variant!r}; known: {list(V)}"); return 2

    doc = docx.Document(_RESUMES / V[variant])

    if args.show:
        idx = index(doc)
        print(f"# {V[variant]}\ntagline: {doc.paragraphs[2].text.strip()}")
        sec = None
        for key in idx["order"]:
            e = idx["entries"][key]
            if e["section"] != sec:
                sec = e["section"]; print(f"\n[{sec}]")
            print(f"  - {key!r}  ({len(e['bullets'])} bullets)")
            for i, b in enumerate(e["bullets"]):
                print(f"      {i}. {b.text[:95]}")
        print("\n[SKILLS]")
        for label, p in idx["skills"].items():
            print(f"  - {label!r}: {p.text.split(':', 1)[1].strip()[:95]}")
        return 0

    if args.worksheet:
        ws = to_worksheet(doc, variant)
        if args.jd:
            jd = Path(args.jd).read_text()
            have = resume_text(ws)
            kws = jd_keywords(jd)
            missing = [k for k in kws if not _present(k, have)]
            claim = json.loads(_CLAIMABLE.read_text()) if _CLAIMABLE.exists() else {}
            ok = set((claim.get("approved") or {}).keys())
            no = set((claim.get("do_not_claim") or {}).keys())
            ws["_ats_present"] = [k for k in kws if _present(k, have)]
            ws["_ats_missing_approved"] = [k for k in missing if k in ok]
            ws["_ats_missing_forbidden"] = [k for k in missing if k in no]
            # everything else: legitimate to add IF experiences.md supports it — a
            # judgement call per application, not a standing permission or ban.
            ws["_ats_missing_unreviewed"] = [k for k in missing
                                             if k not in ok and k not in no]
        out = Path(args.worksheet)
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps(ws, ensure_ascii=False, indent=2))
        n = sum(len(e["bullets"]) for s in ws["sections"] for e in s["entries"])
        print(f"[worksheet] {out}   {len(ws['sections'])} sections, {n} bullets")
        if args.jd:
            print(f"[jd] ATS terms in JD : {len(kws)}")
            print(f"[jd] already covered : {', '.join(ws['_ats_present'])}")
            print(f"[jd] missing, pre-approved    : {', '.join(ws['_ats_missing_approved']) or '-'}")
            print(f"[jd] missing, judge vs experiences.md : {', '.join(ws['_ats_missing_unreviewed']) or '-'}")
            print(f"[jd] missing, ruled out       : {', '.join(ws['_ats_missing_forbidden']) or '-'}")
        return 0

    if not args.apply:
        print("nothing to do: pass --show / --worksheet / --apply / --list"); return 2
    if not args.out:
        print("--out required with --apply"); return 2

    ws = json.loads(Path(args.apply).read_text())
    log = apply_worksheet(doc, ws)

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    for line in log or ["(no changes vs template)"]:
        print("  ·", line)
    print(f"[docx] {out}")

    if args.pdf:
        pdf = to_pdf(out)
        n = page_count(pdf)
        print(f"[pdf ] {pdf}  ({n} page{'s' if n != 1 else ''})"
              + ("   ⚠ NOT ONE PAGE — trim bullets" if n != 1 else ""))

    if args.save_template:
        if not args.reason:
            print("refusing: --save-template needs --reason (rulebook 3.1a: "
                  "template changes must be logged and confirmed with the user)")
            return 2
        name = args.save_template
        fname = f"{_STEM}_{name}.docx"
        (_RESUMES / fname).write_bytes(out.read_bytes())
        # Promote the PDF too. The PDF is what gets uploaded to forms, so
        # leaving a stale one beside a fresh .docx means applying with the old
        # resume — silently, and only visible to whoever is reading it.
        pdf_src = out.with_suffix(".pdf")
        pdf_dst = _RESUMES / f"{_STEM}_{name}.pdf"
        if pdf_src.exists():
            pdf_dst.write_bytes(pdf_src.read_bytes())
            print(f"[template] pdf   {pdf_dst.name}")
        elif pdf_dst.exists():
            pdf_dst.unlink()
            print(f"[template] removed stale {pdf_dst.name} — re-run with --pdf")
        reg = json.loads(_REGISTRY.read_text()) if _REGISTRY.exists() else {}
        reg[name] = fname
        _REGISTRY.write_text(json.dumps(reg, ensure_ascii=False, indent=2))
        chlog = _RESUMES / "TEMPLATE_CHANGELOG.md"
        head = "" if chlog.exists() else "# Template changelog\n\nEvery write to `resumes/` lands here (rulebook 3.1a).\n"
        stamp = __import__("datetime").date.today().isoformat()
        with chlog.open("a") as f:
            f.write(f"{head}\n- {stamp} — `{name}` ({fname}), from `{variant}` "
                    f"via `{args.apply}`\n  reason: {args.reason}\n")
        print(f"[template] promoted to resumes/{fname}  (variant {name!r})")
        print(f"[log ] appended to {chlog}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
